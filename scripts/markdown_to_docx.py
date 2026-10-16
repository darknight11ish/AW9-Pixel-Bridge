#!/usr/bin/env python3
"""Convert an engagement-report Markdown file into its `.docx` companion.

The `.md` file is the source of truth; the `.docx` is a convenience rendering. Keeping
one script means the two cannot drift, which is the failure mode that produced the byte
totals corrected in §12.3 of the report.

Only the Python standard library plus `python-docx` are required.

Usage:
    python scripts/markdown_to_docx.py docs/REPORT.md docs/REPORT.docx
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET_PATTERN = re.compile(r"^(\s*)[-*]\s+(.*)$")
NUMBERED_PATTERN = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")
TABLE_SEPARATOR = re.compile(r"^\|[\s:|-]+\|$")

CODE_FONT = "Consolas"
MONO_COLOR = RGBColor(0x1F, 0x23, 0x28)


def add_inline(paragraph, text: str) -> None:
    """Adds `text` to `paragraph`, rendering **bold** and `code` spans."""
    for token in INLINE_PATTERN.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(9)
            run.font.color.rgb = MONO_COLOR
        else:
            paragraph.add_run(token)


def split_table_row(line: str) -> list[str]:
    """Splits a Markdown table row, tolerating an optional leading/trailing pipe."""
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def render_table(document: Document, rows: list[list[str]]) -> None:
    """Renders a Markdown table as a Word table with a bold header row."""
    columns = max(len(r) for r in rows)
    table = document.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for position in range(columns):
            text = row[position] if position < len(row) else ""
            cell = cells[position]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if index == 0:
                    run.bold = True
    document.add_paragraph()


def convert(markdown_path: Path, docx_path: Path) -> tuple[int, int]:
    """Writes `docx_path` from `markdown_path`; returns (paragraphs, tables)."""
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    document = Document()

    # Body spacing is tightened so the tables dominate the page rather than the text.
    normal = document.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        # Fenced code block: verbatim, monospaced, no inline parsing.
        if stripped.startswith("```"):
            index += 1
            block: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(14)
            paragraph.paragraph_format.space_after = Pt(10)
            run = paragraph.add_run("\n".join(block))
            run.font.name = CODE_FONT
            run.font.size = Pt(8.5)
            run.font.color.rgb = MONO_COLOR
            continue

        # Table: a header row followed by a separator row.
        if stripped.startswith("|") and index + 1 < len(lines) and TABLE_SEPARATOR.match(
            lines[index + 1].strip()
        ):
            rows = [split_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            # A separator may appear mid-table in Markdown; drop any that slipped in.
            rows = [r for r in rows if not all(set(c) <= set(":- ") for c in r)]
            render_table(document, rows)
            continue

        heading = HEADING_PATTERN.match(stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                document.add_heading(text, level=0)
            else:
                document.add_heading(text, level=min(level - 1, 4))
            index += 1
            continue

        if stripped in {"---", "***", "___"}:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run("• • •")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            index += 1
            continue

        # Block quote.
        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            add_inline(paragraph, stripped.lstrip("> ").strip())
            for run in paragraph.runs:
                run.italic = True
            index += 1
            continue

        numbered = NUMBERED_PATTERN.match(line)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(3))
            index += 1
            continue

        bullet = BULLET_PATTERN.match(line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, bullet.group(2))
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        paragraph = document.add_paragraph()
        add_inline(paragraph, stripped)
        index += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(docx_path))
    return len(document.paragraphs), len(document.tables)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("markdown", help="source Markdown file")
    parser.add_argument("docx", help="destination .docx file")
    args = parser.parse_args(argv)

    markdown_path = Path(args.markdown)
    if not markdown_path.is_file():
        raise SystemExit("no such Markdown file: {}".format(markdown_path))

    paragraphs, tables = convert(markdown_path, Path(args.docx))
    print(
        "wrote {} ({} paragraphs, {} tables)".format(args.docx, paragraphs, tables)
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
